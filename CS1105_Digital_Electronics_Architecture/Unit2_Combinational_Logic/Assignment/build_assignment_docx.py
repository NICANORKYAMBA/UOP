#!/usr/bin/env python3
"""
Generate the APA .docx for CS 1105 Unit 2 Assignment (Smart Home Security System).
Double-spaced Times New Roman, native Title/Heading styles, block paragraphs,
monospaced signal-path diagram, Logisim screenshot placeholder, APA references.

Usage:  python3 build_assignment_docx.py
Output: Unit2_Assignment_Activity.docx
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING

FONT = "Times New Roman"
MONO = "Courier New"
OUTPUT = "Unit2_Assignment_Activity.docx"

TITLE = "Smart Home Electronic Security System Using Combinational Circuits"
AUTHOR = "Nicanor Maswili"
AFFIL = "Department of Computer Science, University of the People"
COURSE = "CS 1105: Digital Electronics & Computer Architecture"
INSTRUCTOR = "Instructor: Muhammad Aligohar Bilal"
DUE = "September 16, 2026"

DIAGRAM = [
    "Keypad/Keycard --> ENCODER --> entered code (binary)",
    "Room buttons -------------------> select bits --> MUX (pick room code)",
    "entered code + stored code -----> XNOR per bit --> AND --> GRANT",
    "GRANT + select bits ------------> DECODER/DEMUX --> unlock selected room",
]

REFS = [
    "Harris, D. M., & Harris, S. L. (2012). *Digital design and computer architecture* "
    "(2nd ed.). Morgan Kaufmann.",
    "Mano, M. M., & Ciletti, M. D. (2018). *Digital design: With an introduction to the "
    "Verilog HDL, VHDL, and SystemVerilog* (6th ed.). Pearson.",
    "Ndjountche, T. (2016). *Digital electronics 1: Combinational logic circuits*. John "
    "Wiley & Sons. https://ebookcentral.proquest.com/",
]


def base(doc):
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(12)
    st.paragraph_format.line_spacing = 2.0
    st.paragraph_format.space_after = Pt(0)
    for h, sz in (("Title", 16), ("Heading 1", 13), ("Heading 2", 12)):
        try:
            s = doc.styles[h]
            s.font.name = FONT
            s.font.size = Pt(sz)
            s.font.bold = True
            s.font.color.rgb = RGBColor(0, 0, 0)
            s.paragraph_format.line_spacing = 2.0
            s.paragraph_format.space_before = Pt(0)
            s.paragraph_format.space_after = Pt(0)
        except KeyError:
            pass
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Inches(1)
        sec.left_margin = sec.right_margin = Inches(1)


def para(doc, text="", bold=False, align=None, block=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 2.0
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_after = Pt(10) if block else Pt(0)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.bold = bold
    r.font.name = FONT
    r.font.size = Pt(12)
    return p


def title_style(doc, text):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 2.0
    r = p.add_run(text)
    r.bold = True
    r.font.name = FONT
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(0, 0, 0)


def head(doc, text, level=2, center=False):
    p = doc.add_heading(level=level)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 2.0
    r = p.add_run(text)
    r.bold = True
    r.font.name = FONT
    r.font.size = Pt(13 if level == 1 else 12)
    r.font.color.rgb = RGBColor(0, 0, 0)


def mono(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.left_indent = Inches(0.3)
    r = p.add_run(text)
    r.font.name = MONO
    r.font.size = Pt(10)


def placeholder(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 2.0
    r = p.add_run(text)
    r.bold = True
    r.font.name = FONT
    r.font.size = Pt(12)


def reference(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 2.0
    pf.left_indent = Inches(0.5)
    pf.first_line_indent = Inches(-0.5)
    pf.space_after = Pt(0)
    for i, seg in enumerate(text.split("*")):
        r = p.add_run(seg)
        r.font.name = FONT
        r.font.size = Pt(12)
        r.italic = (i % 2 == 1)


def build():
    doc = Document()
    base(doc)

    for _ in range(4):
        para(doc, "", block=False)
    title_style(doc, TITLE)
    para(doc, "", block=False)
    para(doc, AUTHOR, align=WD_ALIGN_PARAGRAPH.CENTER, block=False)
    para(doc, AFFIL, align=WD_ALIGN_PARAGRAPH.CENTER, block=False)
    para(doc, COURSE, align=WD_ALIGN_PARAGRAPH.CENTER, block=False)
    para(doc, INSTRUCTOR, align=WD_ALIGN_PARAGRAPH.CENTER, block=False)
    para(doc, DUE, align=WD_ALIGN_PARAGRAPH.CENTER, block=False)
    doc.add_page_break()

    head(doc, TITLE, level=1, center=True)

    head(doc, "Introduction")
    para(doc, "As an intern, I have been asked to design a simple electronic security system "
              "for a smart home using combinational circuits. The system controls access to "
              "several rooms and lets authorized users enter by typing a code or presenting a "
              "keycard. Because the decision to grant or deny access depends only on the "
              "current inputs (the entered code and the selected room) and must be immediate, "
              "a combinational design built from logic gates, an encoder, a multiplexer, a "
              "decoder, and a demultiplexer is well suited to the task (Ndjountche, 2016). "
              "This journal presents the design, explains how each component is integrated, "
              "and shows how a demultiplexer extends the system to handle many rooms.")

    head(doc, "(a) Design of the Electronic Security System")
    para(doc, "The system has four stages that flow from input to action:")
    para(doc, "1. Input capture with an encoder. A user enters a code on a keypad or presents "
              "a keycard. The keypad has one line per key; an encoder compresses those "
              "one-hot lines into a compact binary code. For example, a decimal-to-BCD "
              "(10-to-4) encoder turns a pressed digit into a 4-bit value, so the rest of the "
              "circuit works with a few bits instead of many wires (Ndjountche, 2016).")
    para(doc, "2. Room selection with a multiplexer. Each room stores its own authorized "
              "code. A multiplexer (MUX) uses the room-select bits as its control lines to "
              "route the correct stored code for the room the user is trying to enter, so it "
              "chooses which room's rule applies right now.")
    para(doc, "3. Authorization decision with logic gates. A comparator built from XNOR and "
              "AND gates checks whether the entered code matches the stored code for the "
              "selected room. Each bit pair is compared with an XNOR gate, which outputs 1 "
              "only when its two bits are equal; this equality behavior is the standard "
              "building block of a digital comparator (Mano & Ciletti, 2018). The XNOR "
              "outputs feed a single AND gate that outputs 1 (ACCESS GRANTED) only when every "
              "bit matches. For a 4-bit code the grant expression is: Grant = (A0 XNOR S0) "
              "AND (A1 XNOR S1) AND (A2 XNOR S2) AND (A3 XNOR S3), where A is the entered code "
              "and S is the stored code. A keycard-present signal can be AND-ed in for "
              "high-security rooms or OR-ed in for convenience doors, and a NOT gate on a "
              "tamper or door-open sensor can force Grant to 0 under unsafe conditions.")
    para(doc, "4. Room activation with a decoder/demultiplexer. Once access is granted, a "
              "decoder takes the room-select bits and activates exactly one output line, "
              "unlocking the corresponding door, so the grant reaches only the intended room.")
    para(doc, "A simplified signal path is:")
    for line in DIAGRAM:
        mono(doc, line)
    para(doc, "")
    para(doc, "A Logisim implementation of the core of this design is shown in Figures 1 and "
              "2. The circuit realizes the two-bit code comparator (XNOR per bit feeding an "
              "AND gate to produce the GRANT signal) together with a two-to-four decoder built "
              "from AND and NOT gates, so that the single GRANT signal is routed to unlock "
              "exactly one of four rooms selected by the room-select bits R1R0. Figure 1 shows "
              "an authorized attempt (entered code equals the stored code) in which GRANT is "
              "asserted and only the selected room unlocks. Figure 2 shows an unauthorized "
              "attempt (entered code differs from the stored code) in which GRANT is 0 and no "
              "room unlocks, demonstrating that the comparator gates the entire system.")
    placeholder(doc, "[ Insert Figure 1: Logisim screenshot - ACCESS GRANTED "
                     "(code matches, one room unlocks) ]")
    para(doc, "Figure 1. Comparator asserts GRANT and the decoder unlocks only the selected "
              "room.", align=WD_ALIGN_PARAGRAPH.CENTER)
    placeholder(doc, "[ Insert Figure 2: Logisim screenshot - ACCESS DENIED "
                     "(wrong code, no room unlocks) ]")
    para(doc, "Figure 2. With a mismatched code, GRANT is 0 and no room unlocks.",
         align=WD_ALIGN_PARAGRAPH.CENTER)

    head(doc, "(b) Integration of Components and How They Work Together")
    para(doc, "Encoder. The encoder is the entry point. Without it, every key or card line "
              "would need its own wire through the whole circuit. By compressing, say, ten "
              "input lines into a 4-bit code, the encoder reduces wiring and lets the "
              "comparison logic operate on a small, fixed number of bits (Ndjountche, 2016). "
              "Its output, the entered binary code, feeds the comparator.")
    para(doc, "Multiplexer. The MUX makes the system scalable across rooms. Each room has a "
              "stored authorized code; the room-select control lines tell the MUX which "
              "stored code to present to the comparator. One comparison circuit is therefore "
              "reused for every room rather than duplicated per room. The MUX output is one of "
              "the two comparator inputs.")
    para(doc, "Logic gates (XNOR, AND, with NOT and OR as needed). The gates make the actual "
              "authorization decision. Bit by bit, an XNOR gate reports whether the entered "
              "bit equals the stored bit; equality on every bit is required, so the XNOR "
              "outputs feed one AND gate whose output is the GRANT signal. NOT gates can "
              "invert a door-open or tamper sensor to block access under unsafe conditions, "
              "and an OR gate can allow either a valid code or a master keycard to open a "
              "low-security room. Being combinational, the result appears immediately for the "
              "current inputs (Harris & Harris, 2012).")
    para(doc, "Decoder. The decoder converts the room-select bits into a one-hot activation: "
              "for n select bits it drives one of 2^n outputs high (Ndjountche, 2016). "
              "Combined (AND-ed) with the GRANT signal, it ensures the unlock pulse reaches "
              "only the selected, authorized room and never another.")
    para(doc, "Working together. The flow is a clean pipeline: the encoder shrinks the raw "
              "input into a code; the MUX selects the relevant room's stored code; the "
              "gate-based comparator decides grant or deny; and the decoder routes that "
              "decision to the correct door. Each block does one job, and their combinational "
              "nature means the whole path settles to the correct outputs as soon as the "
              "inputs are stable, which is exactly what a responsive door-access system needs.")

    head(doc, "(c) Enhancing the Design with a Demultiplexer")
    para(doc, "A demultiplexer (DEMUX) is the natural way to scale access control to many "
              "rooms efficiently. A DEMUX takes a single data input and, using select lines, "
              "routes it to exactly one of several outputs, the reverse of a multiplexer "
              "(Ndjountche, 2016). Here, the single GRANT signal from the comparator becomes "
              "the DEMUX data input, and the room-select bits become the DEMUX select lines. "
              "The DEMUX then sends the unlock pulse to only the selected room's actuator, "
              "leaving all other doors locked.")
    para(doc, "This is efficient because it reuses one comparison and authorization circuit "
              "for the whole house instead of building a separate authorizer per room: the "
              "MUX chooses which room's code to check on the way in, and the DEMUX distributes "
              "the single decision to the right door on the way out. Adding more rooms simply "
              "requires wider select lines (n select lines support 2^n rooms) rather than "
              "duplicating logic, keeping the design compact and scalable.")

    head(doc, "Academic Integrity Statement")
    para(doc, "This assignment is my own original work. The system design and all "
              "explanations were written by me, and ideas drawn from the course readings are "
              "cited in APA style in the References section below.")

    doc.add_page_break()
    head(doc, "References", level=1, center=True)
    for r in REFS:
        reference(doc, r)

    doc.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    build()
