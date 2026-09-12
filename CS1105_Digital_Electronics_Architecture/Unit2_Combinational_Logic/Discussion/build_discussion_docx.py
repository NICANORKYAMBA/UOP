#!/usr/bin/env python3
"""
Generate the APA .docx for CS 1105 Unit 2 Discussion (Combinational Circuits &
the Elevator Control System). Double-spaced Times New Roman, native Title/Heading
styles, block paragraphs (no first-line indent), hanging-indent APA references.

Usage:  python3 build_discussion_docx.py
Output: Unit2_Discussion_Post.docx
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

FONT = "Times New Roman"
OUTPUT = "Unit2_Discussion_Post.docx"

TITLE = "Combinational Circuits and the Elevator Control System"
AUTHOR = "Nicanor Maswili"
AFFIL = "Department of Computer Science, University of the People"
COURSE = "CS 1105: Digital Electronics & Computer Architecture"
INSTRUCTOR = "Instructor: Muhammad Aligohar Bilal"
DUE = "September 16, 2026"

REFS = [
    "Mano, M. M., & Ciletti, M. D. (2018). *Digital design: With an introduction to the "
    "Verilog HDL, VHDL, and SystemVerilog* (6th ed.). Pearson.",
    "Ndjountche, T. (2016). *Digital electronics 1: Combinational logic circuits*. John "
    "Wiley & Sons. https://ebookcentral.proquest.com/",
]


def base(doc):
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(12)
    st.paragraph_format.line_spacing = 2.0
    st.paragraph_format.space_after = Pt(0)
    for h, sz in (("Title", 16), ("Heading 1", 13), ("Heading 2", 12)):
        try:
            s = doc.styles[h]
            s.font.name = FONT
            s.font.size = Pt(sz)
            s.font.bold = True
            s.font.color.rgb = RGBColor(0, 0, 0)
            s.paragraph_format.line_spacing = 2.0
            s.paragraph_format.space_before = Pt(0)
            s.paragraph_format.space_after = Pt(0)
        except KeyError:
            pass
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Inches(1)
        sec.left_margin = sec.right_margin = Inches(1)


def para(doc, text="", bold=False, align=None, block=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 2.0
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_after = Pt(10) if block else Pt(0)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.bold = bold
    r.font.name = FONT
    r.font.size = Pt(12)
    return p


def title_style(doc, text):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 2.0
    r = p.add_run(text)
    r.bold = True
    r.font.name = FONT
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(0, 0, 0)


def head(doc, text, level=2, center=False):
    p = doc.add_heading(level=level)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 2.0
    r = p.add_run(text)
    r.bold = True
    r.font.name = FONT
    r.font.size = Pt(13 if level == 1 else 12)
    r.font.color.rgb = RGBColor(0, 0, 0)


def reference(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 2.0
    pf.left_indent = Inches(0.5)
    pf.first_line_indent = Inches(-0.5)
    pf.space_after = Pt(0)
    for i, seg in enumerate(text.split("*")):
        r = p.add_run(seg)
        r.font.name = FONT
        r.font.size = Pt(12)
        r.italic = (i % 2 == 1)


def build():
    doc = Document()
    base(doc)

    for _ in range(4):
        para(doc, "", block=False)
    title_style(doc, TITLE)
    para(doc, "", block=False)
    para(doc, AUTHOR, align=WD_ALIGN_PARAGRAPH.CENTER, block=False)
    para(doc, AFFIL, align=WD_ALIGN_PARAGRAPH.CENTER, block=False)
    para(doc, COURSE, align=WD_ALIGN_PARAGRAPH.CENTER, block=False)
    para(doc, INSTRUCTOR, align=WD_ALIGN_PARAGRAPH.CENTER, block=False)
    para(doc, DUE, align=WD_ALIGN_PARAGRAPH.CENTER, block=False)
    doc.add_page_break()

    head(doc, TITLE, level=1, center=True)

    para(doc, "Combinational circuits are the instant decision-makers of digital "
              "electronics: their outputs depend only on the present inputs, produced by "
              "networks of logic gates with no stored state (Ndjountche, 2016). Using the "
              "elevator scenario, this post explains how they differ from sequential "
              "circuits, why they fit elevator control, how a passenger might feel about an "
              "AI-assisted elevator, and where else rapid combinational decision-making adds "
              "value.")

    head(doc, "How Combinational Circuits Differ from Sequential Circuits")
    para(doc, "The defining difference is memory. In a combinational circuit, the output at "
              "any instant is a pure function of the current input combination, evaluated "
              "through gates such as AND, OR, NOT, and XOR; the same inputs always yield the "
              "same outputs, independent of history (Ndjountche, 2016). A sequential circuit "
              "adds memory elements (flip-flops or latches) and is usually driven by a clock, "
              "so its output depends on both the present inputs and the stored past state "
              "(Mano & Ciletti, 2018). Put simply, a combinational circuit answers, given "
              "these inputs right now, what is the output, while a sequential circuit also "
              "asks, and what state were we already in.")
    para(doc, "Two consequences follow: combinational logic has no clock, so it responds as "
              "fast as the gates settle (limited only by propagation delay), and with no "
              "feedback its behavior is fully described by a truth table, making it easy to "
              "verify (Mano & Ciletti, 2018).")
    para(doc, "For an elevator, this stateless speed suits the moment-to-moment safety "
              "decisions. When a floor button is pressed and a sensor reports the car is "
              "elsewhere, a small network of gates can immediately assert the correct "
              "motor-direction and door signals. For example, a simplified interlock such as "
              "MoveUp = RequestAbove AND (NOT DoorOpen) AND (NOT Overweight) is a purely "
              "combinational expression: the gates react instantly to the button and sensor "
              "inputs, guaranteeing the car never moves with an open door or an overloaded "
              "cabin. A complete elevator still needs sequential logic to remember the queue "
              "of requested floors, but the safety-critical decisions about whether to move "
              "this instant are naturally combinational because they must be fast, "
              "predictable, and depend only on the present sensor readings (Ndjountche, 2016).")

    head(doc, "Stepping into an AI-Powered Combinational Elevator")
    para(doc, "As a passenger, I would feel mostly reassured, with a measure of healthy "
              "caution. The exciting possibilities are genuine: an AI layer could learn "
              "building traffic patterns and pre-position the car so waits shrink at rush "
              "hour, group riders bound for nearby floors, and fuse multiple sensor readings "
              "to detect overloading or an obstructed door faster than a human operator. "
              "Because the underlying interlocks remain combinational, the safety response "
              "stays fast and deterministic even while the AI optimizes scheduling above it.")
    para(doc, "My concerns center on trust and failure modes. I would want the AI to optimize "
              "only convenience (routing and timing) while the safety interlocks stay in "
              "fixed, testable combinational logic that the learning model cannot override. I "
              "would also expect a defined fallback if a sensor fails or the AI issues an "
              "unexpected command. With that separation in place, I would step in "
              "confidently, because the guarantees that matter most are enforced by "
              "transparent gate logic rather than by an opaque model.")

    head(doc, "Another Real-World Scenario: Traffic-Intersection Control")
    para(doc, "Beyond elevators, an intelligent traffic-intersection controller benefits "
              "greatly from rapid combinational decision-making. Inductive-loop or camera "
              "sensors report which lanes hold waiting vehicles, whether an emergency vehicle "
              "is approaching, and whether a pedestrian button is pressed. Combinational "
              "logic can instantly derive safe light states; for instance, GrantGreenNorth = "
              "NorthDemand AND (NOT EmergencyCross) AND (NOT ConflictingGreen) ensures two "
              "conflicting directions are never green together, providing an immediate, "
              "verifiable safety guarantee, while a higher-level (sequential or AI) timer "
              "manages the phase sequence. The same outputs-follow-inputs-instantly pattern "
              "powers the arithmetic-logic unit (ALU) inside every processor, where "
              "multiplexers and decoders route and select operands within a single cycle "
              "(Mano & Ciletti, 2018), as well as vending machines and automated warehouse "
              "lifts.")

    head(doc, "Conclusion and Question")
    para(doc, "Combinational circuits excel wherever a system must respond to the present "
              "situation instantly and predictably, which is exactly what safety interlocks "
              "in elevators and traffic systems demand. Pairing them with sequential or AI "
              "layers for memory and optimization yields both speed and intelligence.")
    para(doc, "Question for the group: In a safety-critical system like an elevator, where "
              "should we draw the line between decisions handled by fixed combinational logic "
              "and those handed to an adaptive AI layer, and how would you test that the "
              "combinational safety logic still overrides the AI when the two disagree?")

    doc.add_page_break()
    head(doc, "References", level=1, center=True)
    for r in REFS:
        reference(doc, r)

    doc.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    build()
