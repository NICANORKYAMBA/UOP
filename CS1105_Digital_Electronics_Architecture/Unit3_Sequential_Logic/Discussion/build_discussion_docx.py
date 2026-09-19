#!/usr/bin/env python3
"""
Generate the APA .docx for CS 1105 Unit 3 Discussion (Arcade Game Circuit
Design: D, T, and JK flip-flops). Double-spaced Times New Roman, native
Title/Heading styles, block paragraphs, hanging-indent APA references.

Usage:  python3 build_discussion_docx.py
Output: Unit3_Discussion_Post.docx
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

FONT = "Times New Roman"
OUTPUT = "Unit3_Discussion_Post.docx"

TITLE = "Arcade Game Circuit Design: Choosing Among D, T, and JK Flip-Flops"
AUTHOR = "Nicanor Maswili"
AFFIL = "Department of Computer Science, University of the People"
COURSE = "CS 1105: Digital Electronics & Computer Architecture"
INSTRUCTOR = "Instructor: Muhammad Aligohar Bilal"
DUE = "September 23, 2026"

REFS = [
    "Ndjountche, T. (2016). *Digital electronics 2: Sequential and arithmetic "
    "logic circuits*. ISTE Ltd/John Wiley & Sons. Retrieved from ProQuest "
    "Ebook Central via the UoPeople LIRN Library. "
    "https://ebookcentral.proquest.com/lib/univ-people-ebooks/",
    "Mano, M. M., & Ciletti, M. D. (2018). *Digital design: With an "
    "introduction to the Verilog HDL, VHDL, and SystemVerilog* (6th ed.). "
    "Pearson.",
]


def base(doc):
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(12)
    st.paragraph_format.line_spacing = 2.0
    st.paragraph_format.space_after = Pt(0)
    for h, sz in (("Title", 16), ("Heading 1", 13), ("Heading 2", 12)):
        try:
            s = doc.styles[h]
            s.font.name = FONT
            s.font.size = Pt(sz)
            s.font.bold = True
            s.font.color.rgb = RGBColor(0, 0, 0)
            s.paragraph_format.line_spacing = 2.0
            s.paragraph_format.space_before = Pt(0)
            s.paragraph_format.space_after = Pt(0)
        except KeyError:
            pass
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Inches(1)
        sec.left_margin = sec.right_margin = Inches(1)


def para(doc, text="", bold=False, align=None, block=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 2.0
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_after = Pt(10) if block else Pt(0)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.bold = bold
    r.font.name = FONT
    r.font.size = Pt(12)
    return p


def title_style(doc, text):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 2.0
    r = p.add_run(text)
    r.bold = True
    r.font.name = FONT
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(0, 0, 0)


def head(doc, text, level=2, center=False):
    p = doc.add_heading(level=level)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 2.0
    r = p.add_run(text)
    r.bold = True
    r.font.name = FONT
    r.font.size = Pt(13 if level == 1 else 12)
    r.font.color.rgb = RGBColor(0, 0, 0)


def reference(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 2.0
    pf.left_indent = Inches(0.5)
    pf.first_line_indent = Inches(-0.5)
    pf.space_after = Pt(0)
    for i, seg in enumerate(text.split("*")):
        r = p.add_run(seg)
        r.font.name = FONT
        r.font.size = Pt(12)
        r.italic = (i % 2 == 1)


def build():
    doc = Document()
    base(doc)

    for _ in range(4):
        para(doc, "", block=False)
    title_style(doc, TITLE)
    para(doc, "", block=False)
    para(doc, AUTHOR, align=WD_ALIGN_PARAGRAPH.CENTER, block=False)
    para(doc, AFFIL, align=WD_ALIGN_PARAGRAPH.CENTER, block=False)
    para(doc, COURSE, align=WD_ALIGN_PARAGRAPH.CENTER, block=False)
    para(doc, INSTRUCTOR, align=WD_ALIGN_PARAGRAPH.CENTER, block=False)
    para(doc, DUE, align=WD_ALIGN_PARAGRAPH.CENTER, block=False)
    doc.add_page_break()

    head(doc, TITLE, level=1, center=True)

    para(doc, "The arcade game's control circuit is a good use of sequential logic: unlike "
              "combinational circuits, sequential circuits use memory elements (flip-flops) "
              "whose outputs depend on both the current inputs and the stored past state, "
              "driven by a clock (Ndjountche, 2016). The score must be remembered and "
              "updated, and the lights and sound must toggle on events, so flip-flops are "
              "the right components.")

    head(doc, "Connecting D Flip-Flops to Build the Score Counter")
    para(doc, "A binary counter for the player's score can be built from D flip-flops, one "
              "per bit of the score. A D flip-flop copies its D input to its output Q on the "
              "active clock edge, so it holds one bit until the next update (Ndjountche, "
              "2016). To count, the stages are connected so each toggles at the right time: "
              "in a simple ripple counter the first flip-flop is fed its own inverted output "
              "(D0 = NOT Q0) so it flips every clock pulse, and each following stage is "
              "clocked by the previous stage's output, so it toggles only when the lower "
              "bits roll over. The outputs Q0, Q1, Q2, ... form the binary score, and each "
              "flip-flop stores one bit: Q0 the least-significant bit (value 1), Q1 the next "
              "(value 2), Q2 the next (value 4), and so on. Four D "
              "flip-flops store scores 0000 to 1111 (0 to 15), and adding flip-flops widens "
              "the range. So the information in each D flip-flop is one weighted binary digit "
              "of the current score, and together they store the whole score as a binary "
              "number that, because the flip-flops are edge-triggered, stays stable between "
              "clock pulses (Ndjountche, 2016).")

    head(doc, "Where T Flip-Flops Control Lights and Sound")
    para(doc, "T (toggle) flip-flops are the natural choice for the flashing lights and "
              "sound effects, because a T flip-flop with T held at 1 flips its output on "
              "every clock pulse, producing a steady on-off pattern (Ndjountche, 2016). I "
              "would use one T flip-flop per light that needs to blink: tying T high and "
              "clocking it from a slow pulse makes the light flash at a regular rate. For an "
              "event-driven effect, T becomes the control - setting T = 1 only when a game "
              "event occurs (say, a bonus is hit) lets that event toggle a light or a "
              "sound-enable line, while T = 0 holds the current state. T flip-flops thus fit "
              "any \"flip this on or off each time something happens\" behavior, which is "
              "exactly what blinking indicators and toggled sound effects need.")

    head(doc, "Using JK Flip-Flops Instead of T Flip-Flops")
    para(doc, "A JK flip-flop is the most flexible of the three: with inputs J and K it can "
              "set (J = 1, K = 0), reset (J = 0, K = 1), hold (J = K = 0), or toggle "
              "(J = K = 1) on the clock edge (Mano & Ciletti, 2018). The key behavioral "
              "difference is that a T flip-flop has one input and can only toggle or hold, "
              "so a single pulse merely inverts its current state, whereas a JK can also "
              "deterministically force the output on or off. In effect the JK is a superset "
              "of the T. To reproduce pure toggling I would tie J = K = T, since J = K = 1 "
              "toggles and J = K = 0 holds, so a JK with joined inputs acts as a T flip-flop "
              "(Ndjountche, 2016). Two circuit changes follow. First, each effect flip-flop "
              "now needs two control lines instead of one, so I would add a little "
              "combinational logic to drive J and K from the game-event signals. Second, I "
              "could use that control for exact states rather than toggles - for example, "
              "J = 1, K = 0 to force a warning light on when time is low, and J = 0, K = 1 "
              "to force all lights off at game over, instead of hoping a toggle lands right "
              "(Mano & Ciletti, 2018). The trade-off is more wiring for finer control.")

    head(doc, "Conclusion and Question")
    para(doc, "In short, D flip-flops build the score counter (one bit each), T flip-flops "
              "toggle the lights and sound, and JK flip-flops add independent set/reset at "
              "the cost of an extra input, so the choice depends on how much control each "
              "part needs.")
    para(doc, "Question for the class: When a counter or effect must reset instantly at game "
              "over, is it better to use a flip-flop's asynchronous clear input or to build "
              "the reset into the synchronous logic, and what are the timing risks of each "
              "approach?")

    para(doc, "Word count: 749", block=False)

    doc.add_page_break()
    head(doc, "References", level=1, center=True)
    for r in REFS:
        reference(doc, r)

    doc.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    build()
