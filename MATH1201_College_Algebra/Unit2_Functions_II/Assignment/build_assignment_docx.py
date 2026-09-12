#!/usr/bin/env python3
"""
Generate the APA 7 .docx for MATH 1201 Unit 2 Written Assignment (Functions - II).

Format: Word document, double-spaced, Times New Roman 12pt, 1-inch margins,
native Word Title/Heading styles, block paragraphs, native tables, APA title page,
and a placeholder for the Task 3 GeoGebra graph.

Usage:  python3 build_assignment_docx.py
Output: Unit2_Assignment_Activity.docx
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING

FONT = "Times New Roman"
OUTPUT = "Unit2_Assignment_Activity.docx"

TITLE = "Written Assignment Unit 2: Functions - II"
AUTHOR = "Nicanor Maswili"
AFFIL = "Department of Mathematics, University of the People"
COURSE = "MATH 1201: College Algebra"
INSTRUCTOR = "Instructor: Chibuike Agu"
DUE = "September 16, 2026"

REFS = [
    "Abramson, J. (2023). *Algebra and trigonometry* (2nd ed.). OpenStax. "
    "https://openstax.org/details/books/algebra-and-trigonometry-2e",
    "Khan Academy. (2013, June 3). *Recognizing features of functions (example 2)* "
    "[Video]. YouTube. https://www.youtube.com/",
    "Mathisfun. (n.d.). *Operations with functions*. "
    "https://www.mathsisfun.com/sets/functions-operations.html",
    "Stitz, C., & Zeager, J. (2013). *College algebra*. Stitz Zeager Open Source "
    "Mathematics. https://www.stitz-zeager.com/szca07042013.pdf",
]


def base(doc):
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(12)
    st.paragraph_format.line_spacing = 2.0
    st.paragraph_format.space_after = Pt(0)
    for h, sz in (("Title", 16), ("Heading 1", 13), ("Heading 2", 12)):
        try:
            s = doc.styles[h]
            s.font.name = FONT
            s.font.size = Pt(sz)
            s.font.bold = True
            s.font.color.rgb = RGBColor(0, 0, 0)
            s.paragraph_format.line_spacing = 2.0
            s.paragraph_format.space_before = Pt(0)
            s.paragraph_format.space_after = Pt(0)
        except KeyError:
            pass
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Inches(1)
        sec.left_margin = sec.right_margin = Inches(1)


def para(doc, text="", bold=False, align=None, block=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 2.0
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_after = Pt(10) if block else Pt(0)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.bold = bold
    r.font.name = FONT
    r.font.size = Pt(12)
    return p


def title_style(doc, text):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 2.0
    r = p.add_run(text)
    r.bold = True
    r.font.name = FONT
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p


def head(doc, text, level=2, center=False):
    p = doc.add_heading(level=level)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 2.0
    r = p.add_run(text)
    r.bold = True
    r.font.name = FONT
    r.font.size = Pt(13 if level == 1 else 12)
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p


def table(doc, rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.rows[i].cells[j]
            cell.text = ""
            rp = cell.paragraphs[0]
            rp.paragraph_format.line_spacing = 1.0
            run = rp.add_run(str(val))
            run.font.name = FONT
            run.font.size = Pt(11)
            if i == 0:
                run.bold = True
    para(doc, "")
    return t


def placeholder(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 2.0
    r = p.add_run(text)
    r.bold = True
    r.font.name = FONT
    r.font.size = Pt(12)


def reference(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 2.0
    pf.left_indent = Inches(0.5)
    pf.first_line_indent = Inches(-0.5)
    pf.space_after = Pt(0)
    for i, seg in enumerate(text.split("*")):
        r = p.add_run(seg)
        r.font.name = FONT
        r.font.size = Pt(12)
        r.italic = (i % 2 == 1)


def build():
    doc = Document()
    base(doc)

    # ---- Title page ----
    for _ in range(4):
        para(doc, "", block=False)
    title_style(doc, TITLE)
    para(doc, "", block=False)
    para(doc, AUTHOR, align=WD_ALIGN_PARAGRAPH.CENTER, block=False)
    para(doc, AFFIL, align=WD_ALIGN_PARAGRAPH.CENTER, block=False)
    para(doc, COURSE, align=WD_ALIGN_PARAGRAPH.CENTER, block=False)
    para(doc, INSTRUCTOR, align=WD_ALIGN_PARAGRAPH.CENTER, block=False)
    para(doc, DUE, align=WD_ALIGN_PARAGRAPH.CENTER, block=False)
    doc.add_page_break()

    head(doc, TITLE, level=1, center=True)

    head(doc, "Introduction")
    para(doc, "This assignment applies four core ideas about functions: combining functions "
              "through algebraic operations and composition, inverting a function, "
              "transforming a function's graph, and testing a function for even or odd "
              "symmetry. Each task is worked step by step with the reasoning shown, and the "
              "required graph is produced with the GeoGebra graphing calculator.")

    # ---- Task 1 ----
    head(doc, "Task 1: Algebraic Operations and Composition of Functions")
    para(doc, "Given f(x) = 2x + 1 and g(x) = 3x + 1. The four required operations are "
              "(f/g)(x), (fg)(x), (f o g)(x), and (g o f)(x).")
    para(doc, "(i) Performing the operations.", bold=True)
    para(doc, "Quotient (f/g)(x) = f(x) / g(x) = (2x + 1) / (3x + 1).")
    para(doc, "Product (fg)(x) = (2x + 1)(3x + 1) = 6x^2 + 2x + 3x + 1 = 6x^2 + 5x + 1 "
              "(using the distributive/FOIL method).")
    para(doc, "Composition (f o g)(x) = f(g(x)) = 2(3x + 1) + 1 = 6x + 2 + 1 = 6x + 3 "
              "(apply g first, then substitute into f).")
    para(doc, "Composition (g o f)(x) = g(f(x)) = 3(2x + 1) + 1 = 6x + 3 + 1 = 6x + 4.")
    para(doc, "Verification at x = 5: f(5) = 11, g(5) = 16. Then (f/g)(5) = 11/16; "
              "(fg)(5) = 176 = 6(25)+5(5)+1; (f o g)(5) = f(16) = 33 = 6(5)+3; "
              "(g o f)(5) = g(11) = 34 = 6(5)+4. All check out.")

    para(doc, "(ii) Are fg, f o g, and g o f equal?", bold=True)
    para(doc, "No, all three are different. First, the product (fg)(x) = 6x^2 + 5x + 1 is "
              "quadratic because it multiplies the two outputs, while the compositions "
              "6x + 3 and 6x + 4 are linear because composition substitutes one function "
              "into the other; a quadratic cannot equal a linear function. Second, "
              "(f o g)(x) = 6x + 3 while (g o f)(x) = 6x + 4, so composition is not "
              "commutative - the order matters (Abramson, 2023). Therefore "
              "fg != f o g != g o f.")

    para(doc, "(iii) Domain and range of each operation.", bold=True)
    table(doc, [
        ["Operation", "Result", "Domain", "Range"],
        ["(f/g)(x)", "(2x+1)/(3x+1)", "x != -1/3", "y != 2/3"],
        ["(fg)(x)", "6x^2 + 5x + 1", "(-inf, inf)", "[-1/24, inf)"],
        ["(f o g)(x)", "6x + 3", "(-inf, inf)", "(-inf, inf)"],
        ["(g o f)(x)", "6x + 4", "(-inf, inf)", "(-inf, inf)"],
    ])
    para(doc, "Domains: a quotient is undefined where its denominator is zero, so for (f/g) "
              "set 3x + 1 = 0, giving x = -1/3; the domain excludes -1/3. The product and the "
              "compositions are polynomials with no denominator or radical, so their domain "
              "is all real numbers (Abramson, 2023).")
    para(doc, "Ranges: (f/g) = (2x+1)/(3x+1) is a rational function whose value approaches "
              "the horizontal asymptote 2/3 as x grows large but never reaches it (setting "
              "the expression equal to 2/3 gives the contradiction 3 = 2), so the range "
              "excludes 2/3. The parabola 6x^2 + 5x + 1 opens upward with vertex at "
              "x = -5/12 and minimum value -1/24, so its range is [-1/24, inf). The lines "
              "6x + 3 and 6x + 4 take every real value, so each range is (-inf, inf).")

    # ---- Task 2 ----
    head(doc, "Task 2: Inverse Function - Greenhouse Climate Control")
    para(doc, "The temperature control function is T(C) = sqrt[(20C + 15) / (15C + 16)], "
              "where T is the greenhouse temperature in degrees Celsius and C is the control "
              "setting on the DC inverter.")
    para(doc, "(i) Finding C as a function of T.", bold=True)
    para(doc, "Finding the inverse reverses input and output, expressing the control setting "
              "C in terms of the temperature T (Abramson, 2023). The steps:")
    para(doc, "1. Write the relationship: T = sqrt[(20C + 15) / (15C + 16)].")
    para(doc, "2. Square both sides: T^2 = (20C + 15) / (15C + 16).")
    para(doc, "3. Multiply both sides by (15C + 16): T^2(15C + 16) = 20C + 15, i.e. "
              "15T^2 C + 16T^2 = 20C + 15.")
    para(doc, "4. Collect the C-terms on one side: 15T^2 C - 20C = 15 - 16T^2.")
    para(doc, "5. Factor out C: C(15T^2 - 20) = 15 - 16T^2.")
    para(doc, "6. Divide to isolate C:  C(T) = (15 - 16T^2) / (15T^2 - 20).")
    para(doc, "Verification: a sample setting C = 2 gives T = sqrt(55/46) = 1.0935; "
              "substituting T^2 = 1.1957 into C(T) returns approximately 2.0, recovering the "
              "original setting, which confirms the inverse.")
    para(doc, "(ii) Practical limitations and considerations.", bold=True)
    para(doc, "The square root requires the ratio (20C + 15)/(15C + 16) to be non-negative, "
              "and a real inverter has minimum and maximum settings while a greenhouse has a "
              "realistic temperature band, so the inverse is only meaningful over these "
              "physical intervals. The inverse C(T) is undefined when 15T^2 - 20 = 0, that is "
              "T = sqrt(4/3) which is about 1.15 degrees in the model, so the operating range "
              "must avoid that value. An inverse is valid only where T(C) is one-to-one, "
              "which holds over the monotonic operating range (Abramson, 2023). Finally, "
              "sensor accuracy, response lag, heat loss, and outdoor weather mean the model "
              "is an approximation, so the setting from the inverse is a target that a "
              "feedback loop must still correct.")

    # ---- Task 3 ----
    head(doc, "Task 3: Transformations of a Function")
    para(doc, "The base function is f(x) = fifth root of x (that is, x^(1/5)). The four "
              "transformations are: fifth root of x, plus 6; fifth root of x, minus 6; fifth "
              "root of (50x); and fifth root of (x/50).")
    para(doc, "(i) Graphs.", bold=True)
    para(doc, "All five functions are plotted together in GeoGebra (screenshot below). "
              "GeoGebra input lines: f(x) = nroot(x, 5); then f(x) + 6; f(x) - 6; "
              "nroot(50 x, 5); nroot(x/50, 5).")
    placeholder(doc, "[ Insert GeoGebra graph of f(x) = fifth root of x and its four "
                     "transformations here ]")
    para(doc, "(ii) Explanation of the four transformations.", bold=True)
    para(doc, "Fifth root of x, plus 6 is a vertical shift up 6 units - adding 6 outside the "
              "root raises every point by 6, keeping the shape. Fifth root of x, minus 6 is a "
              "vertical shift down 6 units. Fifth root of (50x) multiplies the input by 50 "
              "inside the root, compressing the graph horizontally toward the y-axis; because "
              "50 = (fifth root of 50)^5, this is also a vertical stretch by about 2.19, so "
              "the curve climbs more steeply. Fifth root of (x/50) divides the input by 50, "
              "stretching the graph horizontally away from the y-axis so it rises more "
              "gradually (equivalently a vertical compression by about 0.46) "
              "(Stitz & Zeager, 2013).")
    para(doc, "(iii) Domain and range observations.", bold=True)
    table(doc, [
        ["Function", "Domain", "Range"],
        ["f(x) = fifth root of x", "(-inf, inf)", "(-inf, inf)"],
        ["fifth root of x + 6", "(-inf, inf)", "(-inf, inf)"],
        ["fifth root of x - 6", "(-inf, inf)", "(-inf, inf)"],
        ["fifth root of (50x)", "(-inf, inf)", "(-inf, inf)"],
        ["fifth root of (x/50)", "(-inf, inf)", "(-inf, inf)"],
    ])
    para(doc, "Because the fifth root is an odd-index radical, it is defined for all real "
              "numbers and outputs all real numbers. Vertical shifts (plus or minus 6) and "
              "horizontal scalings (times 50, divided by 50) change the position and "
              "steepness of the graph but not the set of inputs or outputs, so all five "
              "functions keep domain (-inf, inf) and range (-inf, inf). This contrasts with "
              "an even-index radical such as the square root, whose domain and range would be "
              "restricted to non-negative values.")

    # ---- Task 4 ----
    head(doc, "Task 4: Even/Odd Function Analysis (Alex's Displacement Function)")
    para(doc, "Alex's displacement function is g(t) = 10t^3 / (12t^2 + 53).")
    para(doc, "(i) Even functions and testing g(t).", bold=True)
    para(doc, "A function is even if g(-t) = g(t) for every t in its domain, which means its "
              "graph is symmetric about the y-axis; a function is odd if g(-t) = -g(t), "
              "meaning its graph is symmetric about the origin (Stitz & Zeager, 2013). To "
              "test Alex's function, replace every t with (-t):")
    para(doc, "g(-t) = 10(-t)^3 / (12(-t)^2 + 53). Since (-t)^3 = -t^3 (odd power flips sign) "
              "and (-t)^2 = t^2 (even power stays positive), this becomes "
              "g(-t) = -10t^3 / (12t^2 + 53) = -g(t).")
    para(doc, "Because g(-t) = -g(t) and not g(t), the displacement function is not even; it "
              "is odd. This makes sense structurally: the numerator 10t^3 is an odd-power "
              "term and the denominator 12t^2 + 53 is even, and odd divided by even is odd. A "
              "numerical check confirms it: g(1) = 10/65 = 0.154 while g(-1) = -0.154 = "
              "-g(1).")
    para(doc, "(ii) Interpreting the graph.", bold=True)
    para(doc, "The provided graph passes through the origin and runs from the lower-left, "
              "through (0, 0), to the upper-right, with the left portion below the x-axis and "
              "the right portion above it. It is not symmetric about the y-axis, so it is not "
              "even. It is symmetric about the origin: rotating the graph 180 degrees about "
              "the origin reproduces the same curve, and for every point (t, g(t)) the point "
              "(-t, -g(t)) also lies on the graph, for example (1, 0.15) and (-1, -0.15). "
              "This origin symmetry is the graphical signature of an odd function, which "
              "matches the algebraic result g(-t) = -g(t). Therefore the graph possesses odd "
              "symmetry, not even symmetry.")

    # ---- References ----
    doc.add_page_break()
    head(doc, "References", level=1, center=True)
    for r in REFS:
        reference(doc, r)

    doc.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    build()
