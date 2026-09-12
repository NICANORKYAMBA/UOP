#!/usr/bin/env python3
"""
Generate the submission .docx for CS 1103 Unit 2 Programming Assignment
(Simple E-commerce System using Java packages).

Reads the four source files directly from src/ so the embedded code always
matches the compiled program. Produces: title page, overview, package
structure, the source of each class, verified sample output, screenshot
placeholders, and references. 1.5 spacing, native Word headings, block
paragraphs, monospaced code.

Usage:  python3 build_assignment_docx.py
Output: Unit2_Programming_Assignment.docx
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING

FONT = "Times New Roman"
CODE_FONT = "Courier New"
OUTPUT = "Unit2_Programming_Assignment.docx"

TITLE = "Programming Assignment Unit 2: Simple E-commerce System"
AUTHOR = "Nicanor Maswili"
AFFIL = "Department of Computer Science, University of the People"
COURSE = "CS 1103: Programming 2"
INSTRUCTOR = "Instructor: Chibuike Agu"
DUE = "September 16, 2026"

SRC = "src"
SOURCE_FILES = [
    ("com/ecommerce/Product.java", "Product.java (package com.ecommerce)"),
    ("com/ecommerce/Customer.java", "Customer.java (package com.ecommerce)"),
    ("com/ecommerce/orders/Order.java",
     "Order.java (package com.ecommerce.orders)"),
    ("Main.java", "Main.java (default package)"),
]

SAMPLE_OUTPUT = """===== Welcome to the Online Store =====

Available products:
  [ID 101] Wireless Mouse - $15.99
  [ID 102] Mechanical Keyboard - $49.50
  [ID 103] USB-C Cable - $8.75
  [ID 104] Laptop Stand - $27.00

Customer [ID 1]: Nicanor

Nicanor added Wireless Mouse to the cart.
Nicanor added Mechanical Keyboard to the cart.
Nicanor added USB-C Cable to the cart.
Nicanor removed USB-C Cable from the cart.

Current cart for Nicanor:
  [ID 101] Wireless Mouse - $15.99
  [ID 102] Mechanical Keyboard - $49.50
Cart total: $65.49

===== Order Summary =====
Order ID: 5001
Customer [ID 1]: Nicanor
Products:
  - [ID 101] Wireless Mouse - $15.99
  - [ID 102] Mechanical Keyboard - $49.50
Order Total: $65.49
Status: PLACED
=========================

Order 5001 status updated to SHIPPED.
Final status of order 5001: SHIPPED

--- Validation demo: ordering with an empty cart ---
Order rejected as expected: Cannot place an order with an empty cart.

===== Thank you for shopping with us! ====="""

REFS = [
    "Eck, D. J. (2022). *Introduction to programming using Java* (Version 9, "
    "JavaFX ed.). Hobart and William Smith Colleges. Licensed under CC BY-NC-SA "
    "4.0. https://math.hws.edu/javanotes/",
    "Samoylov, N. (2018). *Introduction to programming: Learn to program in Java "
    "with data structures, algorithms, and logic*. Packt Publishing.",
]


def base(doc):
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(12)
    st.paragraph_format.line_spacing = 1.5
    st.paragraph_format.space_after = Pt(0)
    for h, sz in (("Title", 16), ("Heading 1", 13), ("Heading 2", 12)):
        try:
            s = doc.styles[h]
            s.font.name = FONT
            s.font.size = Pt(sz)
            s.font.bold = True
            s.font.color.rgb = RGBColor(0, 0, 0)
            s.paragraph_format.line_spacing = 1.5
            s.paragraph_format.space_before = Pt(6)
            s.paragraph_format.space_after = Pt(6)
        except KeyError:
            pass
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Inches(1)
        sec.left_margin = sec.right_margin = Inches(1)


def para(doc, text="", bold=False, align=None, block=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_after = Pt(10) if block else Pt(0)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.bold = bold
    r.font.name = FONT
    r.font.size = Pt(12)
    return p


def title_style(doc, text):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text)
    r.bold = True
    r.font.name = FONT
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(0, 0, 0)


def head(doc, text, level=2, center=False):
    p = doc.add_heading(level=level)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text)
    r.bold = True
    r.font.name = FONT
    r.font.size = Pt(13 if level == 1 else 12)
    r.font.color.rgb = RGBColor(0, 0, 0)


def code_block(doc, text):
    for line in text.split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Inches(0.2)
        r = p.add_run(line if line else "")
        r.font.name = CODE_FONT
        r.font.size = Pt(9)


def placeholder(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text)
    r.bold = True
    r.font.name = FONT
    r.font.size = Pt(12)


def reference(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.left_indent = Inches(0.5)
    pf.first_line_indent = Inches(-0.5)
    pf.space_after = Pt(6)
    for i, seg in enumerate(text.split("*")):
        r = p.add_run(seg)
        r.font.name = FONT
        r.font.size = Pt(12)
        r.italic = (i % 2 == 1)


def build():
    doc = Document()
    base(doc)

    # Title page
    for _ in range(4):
        para(doc, "", block=False)
    title_style(doc, TITLE)
    para(doc, "", block=False)
    para(doc, AUTHOR, align=WD_ALIGN_PARAGRAPH.CENTER, block=False)
    para(doc, AFFIL, align=WD_ALIGN_PARAGRAPH.CENTER, block=False)
    para(doc, COURSE, align=WD_ALIGN_PARAGRAPH.CENTER, block=False)
    para(doc, INSTRUCTOR, align=WD_ALIGN_PARAGRAPH.CENTER, block=False)
    para(doc, DUE, align=WD_ALIGN_PARAGRAPH.CENTER, block=False)
    doc.add_page_break()

    head(doc, TITLE, level=1, center=True)

    head(doc, "Overview")
    para(doc, "This program implements a simple e-commerce system in Java, organized with "
              "packages for clean encapsulation. Products and customers live in the "
              "com.ecommerce package, and orders live in the com.ecommerce.orders package, "
              "which imports the Product and Customer classes to show cross-package use of "
              "the import statement. A separate Main class (outside the packages) imports "
              "all three classes and demonstrates the full workflow: browsing a product "
              "catalog, adding and removing items from a shopping cart, calculating the cart "
              "total, placing an order, generating an order summary, and updating the order "
              "status. All fields are private with getters and setters for encapsulation, "
              "and constructors validate their inputs so invalid data is rejected.")

    head(doc, "Package Structure")
    code_block(doc,
               "src/\n"
               "  Main.java                         (default package)\n"
               "  com/\n"
               "    ecommerce/\n"
               "      Product.java                  package com.ecommerce\n"
               "      Customer.java                 package com.ecommerce\n"
               "      orders/\n"
               "        Order.java                  package com.ecommerce.orders")
    para(doc, "")
    para(doc, "Compile and run (from the src directory):")
    code_block(doc,
               "javac -d ../bin com/ecommerce/Product.java "
               "com/ecommerce/Customer.java \\\n"
               "      com/ecommerce/orders/Order.java Main.java\n"
               "java -cp ../bin Main")

    head(doc, "Source Code")
    for path, label in SOURCE_FILES:
        head(doc, label, level=2)
        with open(f"{SRC}/{path}", "r") as f:
            code_block(doc, f.read().rstrip("\n"))
        para(doc, "")

    head(doc, "Sample Run and Output")
    para(doc, "The program was compiled with javac -Xlint:all (no errors or warnings) and "
              "produced the following output. The cart total of $65.49 equals $15.99 + "
              "$49.50 after the USB-C Cable was removed, and the empty-cart order is "
              "correctly rejected by input validation.")
    code_block(doc, SAMPLE_OUTPUT)
    para(doc, "")

    head(doc, "Screenshots of Output")
    para(doc, "Screenshot 1 - Program compiling and running with full output:", bold=True)
    placeholder(doc, "[ Insert Screenshot 1 here ]")
    para(doc, "")
    para(doc, "Screenshot 2 - Input validation (empty-cart order rejected):", bold=True)
    placeholder(doc, "[ Insert Screenshot 2 here ]")

    doc.add_page_break()
    head(doc, "References", level=1, center=True)
    for r in REFS:
        reference(doc, r)

    doc.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    build()
