#!/usr/bin/env python3
"""
Generate the submission .docx for CS 1103 Unit 3 Programming Assignment
(Simple Clock Application using Java threads).

Reads the four source files directly from src/ so the embedded code always
matches the compiled program. Produces: title page, overview, design notes,
the source of each class, verified sample output, screenshot placeholders, and
references. 1.5 spacing, native Word headings, block paragraphs, monospaced
code.

Usage:  python3 build_assignment_docx.py
Output: Unit3_Programming_Assignment.docx
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING

FONT = "Times New Roman"
CODE_FONT = "Courier New"
OUTPUT = "Unit3_Programming_Assignment.docx"

TITLE = "Programming Assignment Unit 3: Simple Clock Application"
AUTHOR = "Nicanor Maswili"
AFFIL = "Department of Computer Science, University of the People"
COURSE = "CS 1103: Programming 2"
INSTRUCTOR = "Instructor: Chibuike Agu"
DUE = "September 23, 2026"

SRC = "src"
SOURCE_FILES = [
    ("Clock.java", "Clock.java"),
    ("ClockUpdater.java", "ClockUpdater.java"),
    ("ClockDisplay.java", "ClockDisplay.java"),
    ("Main.java", "Main.java"),
]

SAMPLE_OUTPUT = """Starting clock application...
Updater thread priority: 1 | Display thread priority: 10
(The clock will run for 15 seconds.)

Current time: 07:30:31 19-09-2026
Current time: 07:30:32 19-09-2026
Current time: 07:30:33 19-09-2026
Current time: 07:30:34 19-09-2026
Current time: 07:30:35 19-09-2026
Current time: 07:30:36 19-09-2026
Current time: 07:30:37 19-09-2026
Current time: 07:30:38 19-09-2026
Current time: 07:30:39 19-09-2026
Current time: 07:30:40 19-09-2026
Current time: 07:30:41 19-09-2026
Current time: 07:30:42 19-09-2026
Current time: 07:30:43 19-09-2026
Current time: 07:30:44 19-09-2026
Current time: 07:30:45 19-09-2026
Clock application stopped."""

REFS = [
    "Eck, D. J. (2022). *Introduction to programming using Java* (Version 9, "
    "JavaFX ed.). Hobart and William Smith Colleges. Licensed under CC BY-NC-SA "
    "4.0. https://math.hws.edu/javanotes/",
    "Samoylov, N. (2018). *Introduction to programming: Learn to program in Java "
    "with data structures, algorithms, and logic*. Packt Publishing.",
]


def base(doc):
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(12)
    st.paragraph_format.line_spacing = 1.5
    st.paragraph_format.space_after = Pt(0)
    for h, sz in (("Title", 16), ("Heading 1", 13), ("Heading 2", 12)):
        try:
            s = doc.styles[h]
            s.font.name = FONT
            s.font.size = Pt(sz)
            s.font.bold = True
            s.font.color.rgb = RGBColor(0, 0, 0)
            s.paragraph_format.line_spacing = 1.5
            s.paragraph_format.space_before = Pt(6)
            s.paragraph_format.space_after = Pt(6)
        except KeyError:
            pass
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Inches(1)
        sec.left_margin = sec.right_margin = Inches(1)


def para(doc, text="", bold=False, align=None, block=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.space_after = Pt(10) if block else Pt(0)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.bold = bold
    r.font.name = FONT
    r.font.size = Pt(12)
    return p


def title_style(doc, text):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text)
    r.bold = True
    r.font.name = FONT
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(0, 0, 0)


def head(doc, text, level=2, center=False):
    p = doc.add_heading(level=level)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text)
    r.bold = True
    r.font.name = FONT
    r.font.size = Pt(13 if level == 1 else 12)
    r.font.color.rgb = RGBColor(0, 0, 0)


def code_block(doc, text):
    for line in text.split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Inches(0.2)
        r = p.add_run(line if line else "")
        r.font.name = CODE_FONT
        r.font.size = Pt(9)


def placeholder(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text)
    r.bold = True
    r.font.name = FONT
    r.font.size = Pt(12)


def reference(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.left_indent = Inches(0.5)
    pf.first_line_indent = Inches(-0.5)
    pf.space_after = Pt(6)
    for i, seg in enumerate(text.split("*")):
        r = p.add_run(seg)
        r.font.name = FONT
        r.font.size = Pt(12)
        r.italic = (i % 2 == 1)


def build():
    doc = Document()
    base(doc)

    # Title page
    for _ in range(4):
        para(doc, "", block=False)
    title_style(doc, TITLE)
    para(doc, "", block=False)
    para(doc, AUTHOR, align=WD_ALIGN_PARAGRAPH.CENTER, block=False)
    para(doc, AFFIL, align=WD_ALIGN_PARAGRAPH.CENTER, block=False)
    para(doc, COURSE, align=WD_ALIGN_PARAGRAPH.CENTER, block=False)
    para(doc, INSTRUCTOR, align=WD_ALIGN_PARAGRAPH.CENTER, block=False)
    para(doc, DUE, align=WD_ALIGN_PARAGRAPH.CENTER, block=False)
    doc.add_page_break()

    head(doc, TITLE, level=1, center=True)

    head(doc, "Overview")
    para(doc, "This program implements a simple clock application in Java that uses "
              "multiple threads to keep and display the current date and time "
              "concurrently. A thread is a separate task that can run at the same time as "
              "other tasks within a single program, which lets the program do more than "
              "one thing at once (Eck, 2022). The application separates two "
              "responsibilities onto two threads: a background thread continuously "
              "updates the clock's stored time, while a separate, higher-priority thread "
              "continuously prints that time to the console. This mirrors the common "
              "pattern of running background work on one thread while a foreground thread "
              "handles output (Samoylov, 2018).")
    para(doc, "The design uses four classes. The Clock class holds the current time and "
              "the logic to format it as HH:mm:ss dd-MM-yyyy. The ClockUpdater and "
              "ClockDisplay classes each implement the Runnable interface and contain the "
              "loop that runs on a thread. The Main class creates one shared clock, wraps "
              "each task in a Thread, assigns thread priorities, starts them, lets the "
              "clock run, and then stops the threads cleanly.")

    head(doc, "Design and Key Decisions")
    para(doc, "Runnable instead of extending Thread. Both tasks implement Runnable rather "
              "than extending the Thread class. Implementing Runnable keeps the task "
              "(what to do) separate from the thread (the worker that runs it), which is "
              "the more flexible approach and is generally preferred (Eck, 2022).")
    para(doc, "Thread priorities. The display thread is given the maximum priority "
              "(Thread.MAX_PRIORITY, value 10) and the background updater thread the "
              "minimum priority (Thread.MIN_PRIORITY, value 1). A higher priority is a "
              "hint to the scheduler to favor that thread, so the display of the time is "
              "prioritized for better timekeeping precision on screen, exactly as the "
              "assignment requires.")
    para(doc, "Safe sharing between threads. The clock's stored time is read by one "
              "thread and written by another, so the updateTime and getFormattedTime "
              "methods are synchronized. This ensures the display thread always reads a "
              "complete, consistent value and never a half-updated one.")
    para(doc, "Clean termination and error handling. Each task uses a volatile boolean "
              "flag and a stop method so its loop can end cleanly. The constructors "
              "validate their arguments and throw IllegalArgumentException for a null "
              "clock or a non-positive interval. Every Thread.sleep call is wrapped in a "
              "try-catch for InterruptedException; on interruption the code restores the "
              "thread's interrupted status and lets the loop exit, which is the "
              "recommended way to handle interruption.")

    head(doc, "Compile and Run")
    para(doc, "From the src directory:")
    code_block(doc,
               "javac -d ../bin Clock.java ClockUpdater.java "
               "ClockDisplay.java Main.java\n"
               "java -cp ../bin Main")

    head(doc, "Source Code")
    for path, label in SOURCE_FILES:
        head(doc, label, level=2)
        with open(f"{SRC}/{path}", "r") as f:
            code_block(doc, f.read().rstrip("\n"))
        para(doc, "")

    head(doc, "Sample Run and Output")
    para(doc, "The program was compiled with javac -Xlint:all (no errors or warnings) and "
              "produced the output below. The header confirms the two thread priorities "
              "(updater 1, display 10), and the clock then updates once per second in the "
              "required HH:mm:ss dd-MM-yyyy format until it stops cleanly after the demo "
              "period. In the live console the time refreshes on a single line, so the "
              "clock reads like a real ticking clock; the lines below show successive "
              "updates.")
    code_block(doc, SAMPLE_OUTPUT)
    para(doc, "")

    head(doc, "Screenshots of Output")
    para(doc, "Screenshot 1 - Program running in IntelliJ, showing the two thread "
              "priorities and the continuously updating clock:", bold=True)
    placeholder(doc, "[ Insert Screenshot 1 here ]")
    para(doc, "")
    para(doc, "Screenshot 2 - The clock a few seconds later, showing the time has "
              "advanced (continuous update):", bold=True)
    placeholder(doc, "[ Insert Screenshot 2 here ]")

    head(doc, "Academic Integrity Statement")
    para(doc, "This assignment is my own original work. I designed and wrote all of the "
              "Java source code and the accompanying explanations myself for this task. "
              "Ideas drawn from the course readings are cited in APA style, and the "
              "sources are listed in the References section below.")

    doc.add_page_break()
    head(doc, "References", level=1, center=True)
    for r in REFS:
        reference(doc, r)

    doc.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    build()
